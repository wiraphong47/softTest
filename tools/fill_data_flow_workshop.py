from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"C:/Users/momil/Documents/ChatGPT/พัฒนาหน้าเว็บไซต์_ในฐานะ Programmer เพื่อลงทะเบียน")
SOURCE = Path(r"C:/Users/momil/Downloads/Workshop2_2_DataFlowDetection.docx")
OUTPUT = ROOT / "output" / "Workshop2_2_DataFlowDetection_Completed.docx"
SCREENSHOT = ROOT / "output" / "workshop_assets" / "website.png"
DFG_IMAGE = ROOT / "output" / "workshop_assets" / "phone_dfg.png"


def format_run(run, size=9.5, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell(cell, text, size=9.5, bold=False, align=None):
    cell.text = ""
    parts = str(text).split("\n")
    for index, part in enumerate(parts):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        if align is not None:
            paragraph.alignment = align
        run = paragraph.add_run(part)
        format_run(run, size, bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_paragraph(paragraph, text, size=10.5):
    paragraph.text = ""
    run = paragraph.add_run(text)
    format_run(run, size)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1


def add_answer_after(document, marker, answer):
    for i, paragraph in enumerate(document.paragraphs):
        if marker in paragraph.text:
            for candidate in document.paragraphs[i + 1:i + 4]:
                if "___" in candidate.text:
                    set_paragraph(candidate, answer)
                    return
    raise ValueError(f"Answer location not found: {marker}")


def generate_dfg_image():
    """Create a compact Mini Data Flow Graph for the logical phone field."""
    canvas = Image.new("RGB", (1200, 520), "white")
    draw = ImageDraw.Draw(canvas)
    font_path = "C:/Windows/Fonts/segoeui.ttf"
    font_bold_path = "C:/Windows/Fonts/segoeuib.ttf"
    title_font = ImageFont.truetype(font_bold_path, 34)
    label_font = ImageFont.truetype(font_bold_path, 24)
    text_font = ImageFont.truetype(font_path, 22)
    small_font = ImageFont.truetype(font_path, 19)

    draw.text((35, 18), "Mini Data Flow Graph: phone", fill="#173a65", font=title_font)
    draw.text((35, 64), "Def = assign value    C-use = display/use value    P-use = condition", fill="#555555", font=small_font)

    nodes = {
        1: (120, 110, "Def(phone)", "createInitialForm\nphone = ''"),
        2: (120, 270, "Def(phone)", "set('phone', x)"),
        3: (120, 425, "C-use(phone)", "value={v.phone}"),
        4: (465, 270, "P-use(phone)", "validateForm:\nempty / format / 10 digits"),
        5: (800, 155, "Error", "errors.phone\nFieldSection / Snackbar"),
        6: (800, 365, "Valid", "no phone error\ncheck next fields"),
    }

    def arrow(start, end, label=""):
        draw.line([start, end], fill="#526d89", width=5)
        dx, dy = end[0] - start[0], end[1] - start[1]
        length = max((dx * dx + dy * dy) ** 0.5, 1)
        ux, uy = dx / length, dy / length
        px, py = -uy, ux
        tip = end
        left = (end[0] - 20 * ux + 10 * px, end[1] - 20 * uy + 10 * py)
        right = (end[0] - 20 * ux - 10 * px, end[1] - 20 * uy - 10 * py)
        draw.polygon([tip, left, right], fill="#526d89")
        if label:
            mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
            draw.text((mx - 8, my - 26), label, fill="#526d89", font=small_font)

    arrow((120, 175), (120, 205))
    arrow((120, 335), (120, 360))
    arrow((185, 270), (395, 270), "submit")
    arrow((530, 250), (730, 165), "invalid")
    arrow((530, 290), (730, 355), "valid")

    for number, (x, y, heading, body) in nodes.items():
        radius = 65
        fill = "#ffe85b" if number <= 4 else ("#f6b4a5" if number == 5 else "#b8e3bf")
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=fill, outline="#30577f", width=4)
        number_box = draw.textbbox((0, 0), str(number), font=label_font)
        draw.text((x - (number_box[2] - number_box[0]) / 2, y - 15), str(number), fill="#173a65", font=label_font)
        draw.text((x + 80, y - 38), heading, fill="#174b87", font=label_font)
        draw.multiline_text((x + 80, y - 5), body, fill="#242424", font=text_font, spacing=3)

    canvas.save(DFG_IMAGE)


def insert_dfg(document):
    for index, paragraph in enumerate(document.paragraphs):
        if "STEP 5 สร้าง Mini Data Flow Graph" in paragraph.text:
            inserted = False
            for candidate in document.paragraphs[index + 1:]:
                if "STEP 6 สร้าง Def-Use Paths" in candidate.text:
                    candidate.paragraph_format.page_break_before = True
                    return
                if "___" in candidate.text:
                    if not inserted:
                        candidate.text = ""
                        candidate.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        candidate.add_run().add_picture(str(DFG_IMAGE), width=Cm(13.1))
                        inserted = True
                    else:
                        candidate.text = ""
    raise ValueError("DFG insertion point not found")


shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)
generate_dfg_image()

# Student and website evidence.
set_cell(doc.tables[0].cell(1, 0), "66040233126", 10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(doc.tables[0].cell(1, 1), "นายวีรพงศ์ วงศ์ชารี", 10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
image_cell = doc.tables[0].cell(1, 2)
image_cell.text = ""
image_paragraph = image_cell.paragraphs[0]
image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_paragraph.add_run().add_picture(str(SCREENSHOT), width=Cm(5.4))
image_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Mission 1: find v.phone.
for paragraph in doc.paragraphs:
    if "ชื่อตัวแปร:" in paragraph.text:
        set_paragraph(paragraph, "ชื่อตัวแปร: v.phone / values.phone    ชนิดข้อมูล: String", 10.5)
        break
add_answer_after(doc, "เหตุผลที่เลือกตัวแปรนี้", "เลือก v.phone เพราะรับค่าจากผู้ใช้โดยตรง และมีเส้นทาง Def → ตรวจสอบความถูกต้อง → แสดงข้อผิดพลาดหรือให้ส่งฟอร์มต่อได้ชัดเจน")

t = doc.tables[5]
rows = [
    ("constants/formOptions.js:8", "createInitialForm กำหนด phone = ''", "Def", "เป็นค่าเริ่มต้นของข้อมูล phone"),
    ("App.jsx:91–92", "set('phone', x) เรียก setV เพื่อกำหนดค่าใหม่ให้ phone", "Def", "ข้อมูลเบอร์โทรจากผู้ใช้ถูกบันทึกใน state"),
    ("App.jsx:332", "TextField อ่าน value={v.phone} เพื่อแสดงในช่องกรอก", "Use", "นำค่าไปแสดงผล (C-use)"),
    ("App.jsx:341", "onChange อ่าน e.target.value แล้วส่งเข้า set('phone', ...)", "Use / Def", "ใช้ค่าที่ผู้ใช้พิมพ์ แล้วกำหนดค่าใหม่ให้ phone"),
    ("validateForm 21–25", "อ่าน values.phone เพื่อตรวจว่าง รูปแบบ และจำนวนตัวเลข", "Use", "ใช้ในเงื่อนไข (P-use) เพื่อเลือกผลลัพธ์"),
]
for row, values in zip(t.rows[1:6], rows):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.7)

t = doc.tables[6]
rows = [
    ("App.jsx:332", "v.phone", "C-use", "ใช้เป็นค่า value ของ TextField เพื่อแสดงเบอร์โทร"),
    ("validateForm 21", "values.phone", "P-use", "ตรวจว่าเบอร์โทรว่างหรือไม่ด้วย if"),
    ("validateForm 22", "values.phone", "P-use", "ตรวจความยาว อักขระที่อนุญาต และการมีตัวเลข"),
    ("validateForm 24", "values.phone", "P-use", "นับตัวเลขแล้วเปรียบเทียบว่าครบ 10 หลักหรือไม่"),
]
for row, values in zip(t.rows[1:5], rows):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.7)
set_cell(t.rows[5].cells[0], "ค่าของ v.phone เดินทางจากช่องกรอกไปยัง state แล้วถูกตรวจด้วยเงื่อนไข ก่อนนำผลไปกำหนดข้อความผิดพลาดและการส่งฟอร์ม", 8.7)

# Mission 2: data map and concepts.
add_answer_after(doc, "วาด Data Map ของตัวแปรที่เลือก", "Data Map: TextField onChange → set('phone', e.target.value) → setV({ ...v, phone: x }) → v.phone → validateForm(v) → errors.phone → FieldSection / Snackbar")
t = doc.tables[8]
set_cell(t.rows[1].cells[1], "ติดตามว่าค่าของตัวแปรถูกกำหนดที่ใด ถูกใช้ที่ใด และเดินทางผ่านเส้นทางใดในโปรแกรม", 9)
set_cell(t.rows[2].cells[1], "ช่วยสร้าง Test Case จากจุดที่ข้อมูลเปลี่ยนค่าและจุดที่ข้อมูลถูกตรวจสอบ จึงพบกรณีขอบเขตหรือเงื่อนไขที่พลาดได้", 9)

t = doc.tables[9]
rows = [
    ("D1: createInitialForm phone = ''", "U1: TextField value={v.phone}", "ใช่", "ค่าเริ่มต้นถูกอ่านเพื่อแสดงช่องกรอก"),
    ("D1: createInitialForm phone = ''", "U2: validateForm ตรวจค่าว่าง", "ใช่", "ไม่มีการกำหนด phone ใหม่ก่อนตรวจ"),
    ("D2: set('phone', x)", "U3: TextField value={v.phone}", "ใช่", "ค่าที่กรอกถูกแสดงกลับในช่องเดิม"),
    ("D2: set('phone', x)", "U4: validateForm ตรวจรูปแบบ", "ใช่", "ค่าจาก input ถูกใช้ในเงื่อนไข validation"),
]
for row, values in zip(t.rows[1:5], rows):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.6)

t = doc.tables[10]
steps = [
    ("1", "สร้าง Data Flow Graph จากโค้ด และระบุโหนด Def, C-use และ P-use"),
    ("2", "หา Def-Use Paths ที่เชื่อมจากการกำหนดค่าไปยังจุดที่ใช้ค่า โดยไม่มีการกำหนดค่าซ้ำระหว่างทาง"),
    ("3", "ออกแบบ Test Case จากแต่ละ Path แล้วเปรียบเทียบ Expected Result กับ Actual Result"),
]
for row, values in zip(t.rows[1:4], steps):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 9)

t = doc.tables[11]
set_cell(t.rows[0].cells[2], "Def → Use ที่เกี่ยวข้อง", 8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
set_cell(t.rows[0].cells[3], "ชนิด Use", 8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
paths = [
    ("Path-1", "เริ่มต้น phone = '' → validateForm ตรวจค่าว่าง", "D1 → U2", "P-use: !values.phone.trim()"),
    ("Path-2", "กรอก 9 หลัก → set phone → นับจำนวนตัวเลข", "D2 → U4", "P-use: จำนวนเลข ≠ 10"),
    ("Path-3", "กรอก 10 หลัก → set phone → validation ผ่าน", "D2 → U4", "P-use ทุกเงื่อนไขเป็น false"),
    ("Path-4", "กรอกอักษรปน → set phone → ตรวจรูปแบบ", "D2 → U4", "P-use: regex ไม่ผ่าน"),
]
for row, values in zip(t.rows[1:5], paths):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.5)

insert_dfg(doc)

# Mission 3: tests and execution records.
t = doc.tables[12]
tests = [
    ("TC-01", "Path-1", "phone = ''", "ไม่ส่งฟอร์มและแสดง ‘กรุณากรอกเบอร์โทร’"),
    ("TC-02", "Path-2", "phone = '081234567'", "ไม่ส่งฟอร์มและแสดง ‘กรุณากรอกเบอร์โทรให้ครบ 10 หลัก’"),
    ("TC-03", "Path-3", "phone = '0812345678'", "ไม่มี error ของ phone และระบบตรวจช่องบังคับถัดไป"),
    ("TC-04", "Path-4", "phone = '08123abc78'", "ไม่ส่งฟอร์มและแจ้งว่าใช้ได้เฉพาะตัวเลข + - และช่องว่าง"),
]
for row, values in zip(t.rows[1:5], tests):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.3)

t = doc.tables[13]
actuals = [
    ("TC-01", "แสดงข้อความตรงตามที่คาด และโฟกัสช่อง Phone", "PASS", "ภาพทดสอบ / npm test"),
    ("TC-02", "แสดงข้อความให้กรอกครบ 10 หลัก และโฟกัสช่อง Phone", "PASS", "ภาพทดสอบ / npm test"),
    ("TC-03", "ไม่มีข้อความผิดพลาดของ Phone แล้วระบบไปตรวจช่อง Date of Birth", "PASS", "npm test"),
    ("TC-04", "แสดงข้อความแจ้งรูปแบบเบอร์โทรไม่ถูกต้อง", "PASS", "npm test"),
]
for row, values in zip(t.rows[1:5], actuals):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.3)

t = doc.tables[14]
set_cell(t.rows[1].cells[1], "Expected และ Actual ตรงกันทั้ง 4 Test Cases: ข้อมูลผิดถูกปฏิเสธ ส่วนข้อมูล 10 หลักผ่าน validation ของ Phone", 9)
set_cell(t.rows[2].cells[1], "ผลการทดสอบอัตโนมัติ 79 กรณี และการทดสอบหน้าเว็บจริงที่แสดงข้อความ/โฟกัสในช่อง Phone", 9)
set_cell(t.rows[3].cells[1], "ไม่เป็น Defect ในเวอร์ชันปัจจุบัน เพราะผลลัพธ์ตรงตามข้อกำหนดเบอร์โทร 10 หลัก", 9)

t = doc.tables[15]
bug_report = [
    ("Bug ID", "N/A - ไม่พบ FAIL ใน Paths ที่เลือก"),
    ("Bug Title", "ไม่มี Bug จากการทดสอบ v.phone ในเวอร์ชันปัจจุบัน"),
    ("Test Case", "TC-01 ถึง TC-04"),
    ("Steps to Reproduce", "ทดสอบข้อมูลว่าง, 9 หลัก, 10 หลัก และตัวอักษรปน ตาม Test Case"),
    ("Expected Result", "ข้อมูลผิดต้องถูกปฏิเสธ ข้อมูล 10 หลักต้องผ่าน validation"),
    ("Actual Result", "ผลลัพธ์ตรงตาม Expected Result ทุกกรณี"),
    ("Severity / Impact", "N/A"),
    ("Evidence", "ผล npm test 79 กรณี และภาพหน้าเว็บไซต์"),
]
for row, values in zip(t.rows[1:9], bug_report):
    set_cell(row.cells[0], values[0], 8.5)
    set_cell(row.cells[1], values[1], 8.5)

t = doc.tables[16]
set_cell(t.rows[1].cells[1], "ไม่ใช้ เพราะไม่พบ Bug ที่ต้องแก้ใน Path ที่เลือก", 9)
set_cell(t.rows[2].cells[1], "N/A", 9)
set_cell(t.rows[3].cells[1], "N/A - ไม่ต้อง Re-test", 9)
set_cell(t.rows[4].cells[1], "ผล TC-01 ถึง TC-04 เป็น PASS", 9)

# Optional case-study comparison table.
t = doc.tables[18]
comparison = [
    ("ตัวแปรที่ติดตาม", "Usage", "v.phone / values.phone"),
    ("Def", "input(Usage)", "createInitialForm และ set('phone', x)"),
    ("Use", "เงื่อนไข if และคำนวณ Bill", "TextField และเงื่อนไข validateForm"),
    ("Def-Use Pair", "Def(Usage) → P-use(Usage)", "Def(v.phone) → P-use(values.phone)"),
    ("Path", "ตามค่าของ Usage ในแต่ละช่วง", "ว่าง / 9 หลัก / 10 หลัก / รูปแบบผิด"),
    ("Test Case", "ค่าขอบเขต Usage", "TC-01 ถึง TC-04"),
]
for row, values in zip(t.rows[1:7], comparison):
    for cell, value in zip(row.cells, values):
        set_cell(cell, value, 8.7)

# Reflection section.
answers = {
    "1. ก่อนทำกิจกรรมนี้": "ก่อนทำกิจกรรมนี้ ฉันคิดว่า Test Case เกิดจากการลองกรอกข้อมูลหลายแบบ แต่ยังไม่ได้เชื่อมโยงกับตัวแปรและเส้นทางใน Source Code อย่างชัดเจน",
    "2. หลังทำกิจกรรมนี้": "หลังทำกิจกรรมนี้ ฉันเห็นว่า Source Code บอกได้ว่าข้อมูลเริ่มต้นและเปลี่ยนค่าอย่างไร ส่วน Test Case สามารถออกแบบจากเส้นทาง Def → Use ของข้อมูลนั้นได้",
    "3. Def และ Use": "Def ช่วยหาแหล่งที่ข้อมูลถูกกำหนดค่า ส่วน Use ช่วยหาจุดที่ข้อมูลถูกนำไปแสดงผลหรือใช้ในเงื่อนไข จึงทำให้เลือก Path ที่ควรทดสอบได้เป็นระบบ",
    "4. จุดไหนของกิจกรรม": "ยังต้องการให้อาจารย์อธิบายเพิ่มเรื่องการวาด DFG ของ React state ที่มีการอัปเดตแบบ asynchronous และวิธีเลือก Def-Use Pair ที่สำคัญเมื่อมีหลายตัวแปร",
    "5. ถ้าต้องกลับไป": "ฉันจะเริ่มจากเลือกตัวแปรสำคัญ เช่น phone, dob หรือ file แล้วติดตาม Def, Use และเงื่อนไข validation ก่อนสร้าง Test Case ทั้งค่าปกติ ค่าขอบเขต และค่าที่ผิดรูปแบบ",
}
for marker, answer in answers.items():
    add_answer_after(doc, marker, answer)

doc.save(OUTPUT)
print(OUTPUT)
